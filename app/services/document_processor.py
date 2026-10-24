"""Turn a file (PDF, DOCX, HTML, Markdown, TXT) into retrievable text chunks.

Two parsers:
- docling (default): layout-aware PDF/DOCX/HTML parsing + token-aware chunking.
  Heavy (downloads layout models on first use) but best on complex PDFs.
- simple: pypdf / python-docx / HTML-to-text + paragraph chunking. Light and fast.

Docling has no plain-text reader, so .txt always uses the simple path (the
original code sent .txt files to docling, where every one of them failed).
"""

from __future__ import annotations

import html
import re
from pathlib import Path

from loguru import logger

from app.config import settings

_SIMPLE_ONLY = {".txt"}
_CHUNK_CHARS = 1_200
_CHUNK_OVERLAP = 150


# kubernetes.io page chrome that survives HTML/PDF conversion. Everything from the
# feedback widget onward is footer, and it appeared in 47 of the 50 shipped docs,
# where it out-ranked real content for many queries.
_FOOTER = re.compile(r"Feedback\s*Was this page helpful\?.*", re.S)
_BOILERPLATE_LINES = re.compile(
    r"(?im)^.*(Thanks for the feedback|answerable question about how to use Kubernetes|"
    r"Open an issue in the GitHub Repository|report a problem|suggest an improvement|"
    r"Last modified .* (AM|PM) [A-Z]{3}|Items on this page refer to third party).*$"
)


def strip_boilerplate(text: str) -> str:
    text = _FOOTER.sub("", text)
    return _BOILERPLATE_LINES.sub("", text)


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _html_to_text(raw: str) -> str:
    raw = re.sub(r"(?is)<(script|style|nav|footer|header)[^>]*>.*?</\1>", " ", raw)
    raw = re.sub(r"(?i)<br\s*/?>|</(p|div|li|h[1-6]|tr|pre)>", "\n", raw)
    raw = re.sub(r"<[^>]+>", " ", raw)
    return html.unescape(raw)


def extract_text_simple(path: Path) -> list[tuple[str, int | None]]:
    """Return [(text, page_number)] using light-weight parsers."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return [(page.extract_text() or "", i + 1) for i, page in enumerate(reader.pages)]
    if suffix == ".docx":
        import docx  # python-docx

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return [("\n".join(parts), None)]
    raw = path.read_text(encoding="utf-8", errors="replace")
    if suffix in (".html", ".htm"):
        raw = _html_to_text(raw)
    return [(raw, None)]


def chunk_text(text: str, size: int = _CHUNK_CHARS, overlap: int = _CHUNK_OVERLAP) -> list[str]:
    """Greedy paragraph packing into ~size-character chunks with a small overlap."""
    text = _clean(strip_boilerplate(text))
    if not text:
        return []
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[str] = []
    current = ""
    for para in paragraphs:
        while len(para) > size:  # split very long paragraphs on sentence/space boundary
            cut = para.rfind(". ", 0, size)
            if cut > size // 2:
                cut += 1
            else:
                cut = para.rfind(" ", 0, size)
            if cut < size // 2:  # no usable boundary (URLs, tables, code): hard cut
                cut = size
            if current:
                chunks.append(current)
                current = ""
            chunks.append(para[:cut].strip())
            # Step back by the overlap, but always move forward at least half a chunk
            # (without this, a short cut could leave `para` unchanged forever).
            para = para[max(cut - overlap, size // 2) :].strip()
        if len(current) + len(para) + 2 <= size:
            current = f"{current}\n\n{para}" if current else para
        else:
            if current:
                chunks.append(current)
            tail = current[-overlap:] if current else ""
            current = f"{tail}\n\n{para}".strip() if tail else para
    if current:
        chunks.append(current)
    return [c for c in chunks if len(c) >= 40]


class DocumentProcessor:
    def __init__(self, parser: str | None = None):
        self.parser = (parser or settings.doc_parser).lower()
        self._converter = None
        self._chunker = None

    def _docling(self):
        if self._converter is None:
            from docling.chunking import HybridChunker
            from docling.datamodel.accelerator_options import AcceleratorDevice, AcceleratorOptions
            from docling.datamodel.base_models import InputFormat
            from docling.datamodel.pipeline_options import PdfPipelineOptions
            from docling.document_converter import DocumentConverter, PdfFormatOption

            pipeline_options = PdfPipelineOptions()
            # AUTO picks CUDA / Apple MPS / CPU. The original hard-coded MPS,
            # which fails on Linux and inside Docker.
            pipeline_options.accelerator_options = AcceleratorOptions(
                num_threads=4, device=AcceleratorDevice.AUTO
            )
            self._converter = DocumentConverter(
                format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=pipeline_options)}
            )
            self._chunker = HybridChunker()
        return self._converter, self._chunker

    def _process_docling(self, file_path: str, source_name: str) -> list[dict]:
        converter, chunker = self._docling()
        doc = converter.convert(file_path).document
        chunks = []
        for chunk in chunker.chunk(doc):
            text = chunker.contextualize(chunk) if hasattr(chunker, "contextualize") else chunk.text
            text = _clean(strip_boilerplate(text))
            if len(text) < 40:
                continue
            meta = {"text": text, "source": source_name}
            if hasattr(chunk, "meta") and hasattr(chunk.meta, "doc_items"):
                items = chunk.meta.doc_items
                if items and hasattr(items[0], "prov") and items[0].prov:
                    meta["page_number"] = items[0].prov[0].page_no
            chunks.append(meta)
        return chunks

    def _process_simple(self, file_path: str, source_name: str) -> list[dict]:
        chunks = []
        for text, page in extract_text_simple(Path(file_path)):
            for piece in chunk_text(text):
                meta: dict = {"text": piece, "source": source_name}
                if page is not None:
                    meta["page_number"] = page
                chunks.append(meta)
        return chunks

    def process_document(self, file_path: str, source_name: str | None = None) -> list[dict]:
        source_name = source_name or Path(file_path).name
        suffix = Path(file_path).suffix.lower()
        use_docling = self.parser == "docling" and suffix not in _SIMPLE_ONLY

        chunks: list[dict] = []
        if use_docling:
            try:
                chunks = self._process_docling(file_path, source_name)
            except Exception as exc:  # noqa: BLE001
                logger.warning("docling failed on {} ({}); using simple parser", source_name, exc)
        if not chunks:
            chunks = self._process_simple(file_path, source_name)

        logger.info("Processed {} chunks from {}", len(chunks), source_name)
        return chunks
